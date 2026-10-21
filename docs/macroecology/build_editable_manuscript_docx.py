"""Build a coauthor-editable manuscript from the source-reconciled FCP companion draft.

The GitHub Markdown files are the canonical authoring source. DOCX is a reader-facing
review copy. No figure preview is represented as final validated biology.
"""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).parent
OUT=ROOT/'FCP_MACROECOLOGY_MANUSCRIPT_WORKING_V01.docx'
DOC=Document()
sec=DOC.sections[0]
sec.top_margin=Inches(.82)
sec.bottom_margin=Inches(.8)
sec.left_margin=Inches(.87)
sec.right_margin=Inches(.83)
style=DOC.styles['Normal']
style.font.name='Times New Roman'
style.font.size=Pt(10.5)
style.paragraph_format.space_after=Pt(5)
style.paragraph_format.line_spacing=1.3
for hn,size in [('Title',17),('Heading 1',13),('Heading 2',11.5)]:
    st=DOC.styles[hn]
    st.font.name='Times New Roman'
    st.font.size=Pt(size)
    st.font.bold=True
    st.paragraph_format.space_before=Pt(12)
    st.paragraph_format.space_after=Pt(6)

hdr=sec.header.paragraphs[0]
hdr.text='FCP macroecology — coauthor review draft, v0.1 (14 Oct 2026)'
hdr.alignment=WD_ALIGN_PARAGRAPH.RIGHT
hdr.style=DOC.styles['Normal']
for run in hdr.runs:
    run.font.size=Pt(8)
    run.italic=True

foot=sec.footer.paragraphs[0]
foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=foot.add_run('Working draft · not submitted · ')
r.font.size=Pt(8)
fld=OxmlElement('w:fldSimple')
fld.set(qn('w:instr'),'PAGE')
foot._p.append(fld)


def format_inline(p, s):
    # Markdown links are rendered as readable labels + an ordinary URL only if a
    # link is not a GitHub in-text crossreference. All external references remain.
    s=re.sub(r'\[([^\]]+)\]\((https?://[^\s)]+)\)',lambda m: m.group(1)+' ('+m.group(2)+')',s)
    parts=re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)',s)
    for piece in parts:
        if not piece:continue
        bold=piece.startswith('**') and piece.endswith('**')
        italic=(not bold) and piece.startswith('*') and piece.endswith('*')
        code=piece.startswith('`') and piece.endswith('`')
        text=piece[2:-2] if bold else piece[1:-1] if (italic or code) else piece
        run=p.add_run(text)
        if bold:run.bold=True
        if italic:run.italic=True
        if code:run.font.name='Consolas';run.font.size=Pt(9)


def add_markdown(path):
    rows=path.read_text().splitlines()
    buffer=[]
    def flush():
        nonlocal buffer
        if not buffer:return
        para=DOC.add_paragraph()
        format_inline(para,' '.join(buffer).strip())
        buffer=[]
    for line in rows:
        t=line.strip()
        if not t:
            flush();continue
        if t=='---':flush();continue
        if t.startswith('# '):
            flush();p=DOC.add_paragraph(style='Title');format_inline(p,t[2:]);continue
        if t.startswith('## '):
            flush();p=DOC.add_paragraph(style='Heading 1');format_inline(p,t[3:]);continue
        if t.startswith('### '):
            flush();p=DOC.add_paragraph(style='Heading 2');format_inline(p,t[4:]);continue
        if t.startswith('- '):
            flush();p=DOC.add_paragraph(style='List Bullet');format_inline(p,t[2:]);continue
        if any(t.startswith(marker) for marker in (
             '**Research article', '**Authors:', '**Affiliations:',
             '**Corresponding author:', '**Status:', '**Suggested scope:', '**Keywords')):
            flush()
            p=DOC.add_paragraph()
            p.paragraph_format.space_after=Pt(2)
            format_inline(p,t)
            continue
        if t.startswith('**') and t.endswith('**') and len(t)<150 and ':' not in t:
            flush();p=DOC.add_paragraph();format_inline(p,t);continue
        buffer.append(t)
    flush()

add_markdown(ROOT/'MANUSCRIPT_V0_1.md')

# Append publication table bodies. These tables are evidence-indexed; they do not
# mix unlike binary NHM source effect sizes into photographic four-colour Table 2.
DOC.add_page_break()
p=DOC.add_paragraph(style='Heading 1');p.add_run('Main Tables — coauthor working copy')
lines=(ROOT/'MAIN_TABLES.md').read_text().splitlines()
sections={}
current=None
for line in lines:
    if line.startswith('## '):
        current=line[3:]
        sections[current]=[]
    elif current is not None:
        sections[current].append(line)

for heading, rows in sections.items():
    p=DOC.add_paragraph(style='Heading 2')
    p.add_run(heading)
    buffer=[]
    done=False
    # The simple parser expects one GitHub-flavour markdown table per section.
    for line in rows:
        if line.startswith('|'):
            buffer.append(line)
        elif buffer:
            if not done:
                done=True
                break
    if not buffer:continue
    tablelines=[l.strip() for l in buffer if not re.match(r'^\|[\s:|\-]+\|$',l.strip())]
    cells=[[s.strip().replace('**','').replace('`','') for s in l.strip().strip('|').split('|')] for l in tablelines]
    if not cells:continue
    nc=len(cells[0]); assert all(len(a)==nc for a in cells)
    table=DOC.add_table(rows=0,cols=nc)
    table.style='Table Grid'
    table.alignment=WD_TABLE_ALIGNMENT.CENTER
    table.autofit=True
    for i,values in enumerate(cells):
        row=table.add_row()
        for cell,text in zip(row.cells,values):
            cell.text=text
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(1.5)
                para.paragraph_format.space_before=Pt(1.5)
                para.paragraph_format.line_spacing=1.05
                for run in para.runs:
                    run.font.name='Times New Roman'
                    run.font.size=Pt(7.2 if nc>=6 else 7.5)
                    if i==0:run.bold=True
        if i==0:
            trPr=row._tr.get_or_add_trPr()
            h=OxmlElement('w:tblHeader')
            h.set(qn('w:val'),'true')
            trPr.append(h)
        else:
            trPr=row._tr.get_or_add_trPr()
            cs=OxmlElement('w:cantSplit')
            trPr.append(cs)
    DOC.add_paragraph('')

for label,imgname in [
    ('Figure 2. Photographic colour directions — prototype, not submission-final',
     'FIG2_macro_four_hue_direction_preview.png'),
    ('Figure 3. Photographic and spatial support — prototype, not submission-final',
     'FIG3_BIO12_spatial_support_preview.png')]:
    DOC.add_page_break()
    p=DOC.add_paragraph(style='Heading 1');p.add_run(label)
    path=ROOT/'figures'/imgname
    if path.exists():
        pp=DOC.add_paragraph()
        pp.alignment=WD_ALIGN_PARAGRAPH.CENTER
        pp.add_run().add_picture(str(path),width=Inches(6.65))
    cap=DOC.add_paragraph('The figure derives from the frozen FCP PR #193/#194 observational source ledgers. Retrospective sampling, within-scope confidence intervals and uncertainty about causal selection remain; see Methods and figure plan.')
    cap.style=DOC.styles['Normal']
    for run in cap.runs:run.font.size=Pt(8.5)

DOC.core_properties.title='Context dependence of flower-colour-climate associations (working draft)'
DOC.core_properties.subject='Global FCP macroecology companion manuscript; not submitted'
DOC.core_properties.author='FCP research manuscript — authors to confirm'
DOC.core_properties.keywords='global flower colour; climate; spatial conditioning; macroecology'
DOC.save(OUT)
print(OUT)
