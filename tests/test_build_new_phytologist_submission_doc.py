import importlib.util
import zipfile
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "scripts" / "publication" / "build_new_phytologist_submission_doc.py"
SOURCE = ROOT / "docs" / "POLYMORPHISM_MANUSCRIPT_NEW_PHYTOLOGIST.md"
FIGURES = ROOT / "docs" / "figures" / "polymorphism_20260918"


def load_module():
    assert SCRIPT.exists(), "New Phytologist submission DOCX builder is not implemented"
    spec = importlib.util.spec_from_file_location("new_phy_submission_doc", SCRIPT)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(module)
    return module


def test_build_review_ready_docx_contract(tmp_path: Path) -> None:
    module = load_module()
    out = tmp_path / "fcp_new_phytologist_submission.docx"
    result = module.build_submission_docx(
        source_path=SOURCE,
        figures_dir=FIGURES,
        output_path=out,
    )

    assert result == out
    assert out.exists() and out.stat().st_size > 300_000

    doc = Document(out)
    joined = "\n".join(p.text for p in doc.paragraphs)
    assert "A recurrent achromatic" in joined
    assert "[AUTHOR LIST TO CONFIRM]" in joined
    assert "[AFFILIATIONS TO INSERT]" in joined
    assert "[NAME / EMAIL TO INSERT]" in joined
    assert "H2_PROSPECTIVE_WHITE_AXIS_CONFIRMED" in joined
    assert "not an independent-source replication" in joined
    assert len(doc.tables) >= 1

    with zipfile.ZipFile(out) as zf:
        names = zf.namelist()
        document_xml = zf.read("word/document.xml").decode("utf-8")
        styles_xml = zf.read("word/styles.xml").decode("utf-8")
        footer_xml = "\n".join(
            zf.read(name).decode("utf-8")
            for name in names
            if name.startswith("word/footer") and name.endswith(".xml")
        )
        media = [name for name in names if name.startswith("word/media/")]

    assert len(media) >= 5
    assert "<w:lnNumType" in document_xml
    assert 'w:restart="continuous"' in document_xml
    assert 'w:countBy="1"' in document_xml
    assert "PAGE" in footer_xml
    assert 'w:line="360"' in styles_xml
