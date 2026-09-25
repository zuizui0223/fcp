#!/usr/bin/env python3
from __future__ import annotations
import argparse,json,re
from pathlib import Path
import numpy as np
import pandas as pd
from docx import Document

NUM=re.compile(r"(?<![A-Za-z])([0-9]+(?:\.[0-9]+)?)")

def parse_frequency(s):
    raw=str(s).strip()
    norm=raw.replace("–","-").replace("—","-").replace("“","").replace("”","")
    lower=norm.lower()
    vals=[float(x) for x in NUM.findall(norm)]
    qualitative=None
    if "extremely rare" in lower:
        qualitative="extremely rare"
    elif re.search(r"\brare\b",lower):
        qualitative="rare"
    censored=norm.lstrip().startswith("<")
    greenhouse="greenhouse" in lower
    if qualitative is not None:
        return {
            "raw_frequency":raw,"numeric":False,"reported_min":None,"reported_max_or_upper":None,
            "censored_upper":False,"qualitative_frequency":qualitative,"greenhouse_flag":greenhouse
        }
    if not vals:
        raise ValueError(f"unparseable frequency: {raw!r}")
    if censored:
        return {
            "raw_frequency":raw,"numeric":True,"reported_min":0.0,"reported_max_or_upper":vals[0],
            "censored_upper":True,"qualitative_frequency":None,"greenhouse_flag":greenhouse
        }
    return {
        "raw_frequency":raw,"numeric":True,"reported_min":min(vals),"reported_max_or_upper":max(vals),
        "censored_upper":False,"qualitative_frequency":None,"greenhouse_flag":greenhouse
    }

def main():
    p=argparse.ArgumentParser()
    p.add_argument("--docx",required=True)
    p.add_argument("--outdir",required=True)
    a=p.parse_args()
    out=Path(a.outdir); out.mkdir(parents=True,exist_ok=True)

    doc=Document(a.docx)
    if len(doc.tables)!=1:
        raise SystemExit(f"expected one table, found {len(doc.tables)}")
    t=doc.tables[0]
    if len(t.rows)!=32 or len(t.columns)!=5:
        raise SystemExit(f"unexpected table dimensions {len(t.rows)} x {len(t.columns)}")

    rows=[[c.text.replace("\n"," | ").strip() for c in r.cells] for r in t.rows]
    if "White-petal phenotype frequency" not in rows[1][1]:
        raise SystemExit(f"unexpected header: {rows[1]}")
    if not rows[3][0].startswith("Species with PAL"):
        raise SystemExit("PAL section header not found")
    if not rows[17][0].startswith("Species with WAL"):
        raise SystemExit("WAL section header not found")

    parsed=[]
    for cls,start,end in [("PAL",4,17),("WAL",18,31)]:
        for i in range(start,end):
            cells=rows[i]
            species=cells[0].strip()
            freq=parse_frequency(cells[1])
            parsed.append({
                "architecture_class":cls,
                "species_source_text":species,
                "nonanthocyanin_flavonoids_flower":cells[2],
                "nonanthocyanin_flavonoids_vegetative":cells[3],
                "reference":cells[4],
                "source_excel_row":i+1,
                **freq,
            })

    d=pd.DataFrame(parsed)
    d.to_csv(out/"cross_species_frequency_registry.csv",index=False)

    pal=d[d.architecture_class.eq("PAL")].copy()
    wal=d[d.architecture_class.eq("WAL")].copy()
    pnum=pal[pal.numeric].copy()
    wnum=wal[wal.numeric].copy()
    if len(pal)!=13 or len(wal)!=13:
        raise SystemExit(f"unexpected class sizes PAL={len(pal)} WAL={len(wal)}")
    if len(pnum)!=13:
        raise SystemExit(f"expected all 13 PAL rows numeric, got {len(pnum)}")

    palmax=pnum.reported_max_or_upper.to_numpy(float)
    walupper=wnum.reported_max_or_upper.to_numpy(float)
    min_pal=float(np.min(palmax))
    max_wal=float(np.max(walupper))
    separated=bool(min_pal>max_wal)

    published=d.loc[d.reference.astype(str).str.contains(r"\\[",regex=True) & ~d.greenhouse_flag.astype(bool)].copy()
    pub_summary={}
    for cls in ["PAL","WAL"]:
        q=published.loc[published.architecture_class.eq(cls)].copy()
        num=q.loc[q.numeric].copy()
        pub_summary[cls]={
            "species":int(len(q)),
            "numeric_species":int(len(num)),
            "qualitative_rare_species":int(q.qualitative_frequency.notna().sum()),
            "numeric_min_upper_percent":float(num.reported_max_or_upper.min()) if len(num) else None,
            "numeric_median_upper_percent":float(num.reported_max_or_upper.median()) if len(num) else None,
            "numeric_max_upper_percent":float(num.reported_max_or_upper.max()) if len(num) else None
        }
    pub_sep=bool(
        pub_summary["PAL"]["numeric_min_upper_percent"] >
        pub_summary["WAL"]["numeric_max_upper_percent"]
    )

    result={
      "schema":"fcp_silene_cross_species_pal_wal_frequency_v1",
      "status":"complete",
      "role":"post_publication_cross_species_descriptive_reproduction",
      "source":"Supplementary Table S1, Del Valle et al. 2019",
      "doi":"10.1186/s12870-019-2082-6",
      "PAL":{
        "species":int(len(pal)),
        "numeric_species":int(len(pnum)),
        "documented_max_frequency_min_percent":min_pal,
        "documented_max_frequency_median_percent":float(np.median(palmax)),
        "documented_max_frequency_max_percent":float(np.max(palmax))
      },
      "WAL":{
        "species":int(len(wal)),
        "numeric_species":int(len(wnum)),
        "qualitative_rare_species":int(wal.qualitative_frequency.notna().sum()),
        "numeric_upper_bound_min_percent":float(np.min(walupper)),
        "numeric_upper_bound_median_percent":float(np.median(walupper)),
        "numeric_upper_bound_max_percent":max_wal,
        "greenhouse_flagged_numeric_species":int(wnum.greenhouse_flag.sum())
      },
      "frequency_scale_separated":separated,
      "conservative_separation_factor_min_PALmax_over_max_WALupper":float(min_pal/max_wal),
      "postopen_published_only_non_greenhouse_sensitivity":{
        "role":"post-open descriptive robustness; cannot alter the primary source-table summary",
        "filter":"reference contains a published numeric citation [..] and greenhouse_flag is false",
        "PAL":pub_summary["PAL"],
        "WAL":pub_summary["WAL"],
        "frequency_scale_separated":pub_sep,
        "conservative_separation_factor":float(
            pub_summary["PAL"]["numeric_min_upper_percent"]/
            pub_summary["WAL"]["numeric_max_upper_percent"]
        )
      },
      "interpretation":[
        "All 13 PAL systems in the source table have a documented maximum white frequency of at least 10%.",
        "Every numerically reported WAL system has an upper frequency bound at or below 1.4%; four additional WAL systems are reported qualitatively as rare or extremely rare.",
        "The reported frequency scales do not overlap under this conservative summary."
      ],
      "hard_nonclaims":[
        "source studies differ in survey design and effort",
        "qualitative rare entries are not converted to numbers",
        "censored <x values are retained as upper bounds",
        "no comparative p-value is reported",
        "post-publication literature synthesis, not independent confirmation"
      ]
    }
    (out/"result.json").write_text(json.dumps(result,indent=2)+"\n")
    print(json.dumps(result,indent=2))

if __name__=="__main__": main()
