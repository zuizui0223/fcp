#!/usr/bin/env python3
from __future__ import annotations
import argparse,json
from pathlib import Path
import pandas as pd
from docx import Document

def main():
    p=argparse.ArgumentParser()
    p.add_argument("--docx",required=True)
    p.add_argument("--outdir",required=True)
    a=p.parse_args()
    out=Path(a.outdir); out.mkdir(parents=True,exist_ok=True)
    doc=Document(a.docx)
    manifest=[]; rows=[]
    for ti,t in enumerate(doc.tables,1):
        manifest.append({"table_index":ti,"rows":len(t.rows),"cols":len(t.columns)})
        for ri,row in enumerate(t.rows,1):
            vals=[c.text.replace("\n"," | ").strip() for c in row.cells]
            rec={"table_index":ti,"row_index":ri}
            for j,v in enumerate(vals,1): rec[f"c{j}"]=v
            rows.append(rec)
    pd.DataFrame(manifest).to_csv(out/"table_manifest.csv",index=False)
    raw=pd.DataFrame(rows); raw.to_csv(out/"all_docx_rows.csv",index=False)
    result={"schema":"fcp_silene_s3_biochemistry_preflight_v1","status":"complete","tables":manifest,"paragraphs":len(doc.paragraphs),"rule":"schema-only; no biochemical phenotype comparison calculated"}
    (out/"result.json").write_text(json.dumps(result,indent=2)+"\n")
    print(json.dumps(result,indent=2))
    print(raw.head(180).to_string(index=False))
if __name__=="__main__": main()
