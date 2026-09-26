#!/usr/bin/env python3
from __future__ import annotations
import argparse, json, math
from pathlib import Path
import numpy as np
import pandas as pd
from docx import Document

YEARS=[2010,2013,2014,2016,2018]
PHEN=["PAL","WAL"]

def parse_num(x):
    s=str(x).strip()
    if s in {"","-","–","—"}:
        return None
    try:
        return float(s)
    except Exception:
        raise ValueError(f"unparseable frequency cell: {x!r}")

def main():
    p=argparse.ArgumentParser()
    p.add_argument("--docx",required=True)
    p.add_argument("--outdir",required=True)
    a=p.parse_args()
    out=Path(a.outdir); out.mkdir(parents=True,exist_ok=True)

    doc=Document(a.docx)
    if len(doc.tables)!=1:
        raise SystemExit(f"expected one table, found {len(doc.tables)}")
    t=doc.tables[0]
    if len(t.rows)!=25 or len(t.columns)!=14:
        raise SystemExit(f"unexpected table dimensions: {len(t.rows)} x {len(t.columns)}")

    h1=[c.text.strip() for c in t.rows[1].cells]
    h2=[c.text.strip() for c in t.rows[2].cells]
    expected=[]
    for y in YEARS:
        expected.extend([(str(y),"PAL"),(str(y),"WAL")])
    actual=[(h1[j],h2[j]) for j in range(4,14)]
    if actual!=expected:
        raise SystemExit(f"unexpected year/phenotype schema: {actual}")

    pops=[]; long=[]
    for row in t.rows[3:24]:
        c=[x.text.strip() for x in row.cells]
        pop=c[0]
        pops.append({
            "population":pop,
            "latitude_text":c[1],
            "longitude_text":c[2],
            "estimated_population_size":c[3],
        })
        for k,(year,phen) in enumerate([(y,p) for y in YEARS for p in PHEN],start=4):
            val=parse_num(c[k])
            if val is None:
                continue
            long.append({
                "population":pop,
                "year":year,
                "phenotype":phen,
                "frequency_percent":val,
            })
    long=pd.DataFrame(long)
    pd.DataFrame(pops).to_csv(out/"populations.csv",index=False)
    long.to_csv(out/"population_year_frequencies.csv",index=False)

    phenotype_results={}
    trajectory_rows=[]
    for phen in PHEN:
        q=long.loc[long.phenotype.eq(phen)].copy()
        qp=q.loc[q.frequency_percent.gt(0)].copy()
        g=(q.groupby("population")
             .agg(observed_years=("year","nunique"),
                  positive_years=("frequency_percent",lambda s:int((s>0).sum())),
                  min_frequency=("frequency_percent","min"),
                  median_frequency=("frequency_percent","median"),
                  max_frequency=("frequency_percent","max"))
             .reset_index())
        for _,r in g.iterrows():
            trajectory_rows.append({"phenotype":phen,**r.to_dict()})
        positive_values=qp.frequency_percent.to_numpy(float)
        phenotype_results[phen]={
            "populations_observed":int(q.population.nunique()),
            "population_year_observations":int(len(q)),
            "populations_ever_positive":int(qp.population.nunique()),
            "positive_population_years":int(len(qp)),
            "positive_frequency_min_percent":float(np.min(positive_values)) if len(positive_values) else None,
            "positive_frequency_median_percent":float(np.median(positive_values)) if len(positive_values) else None,
            "positive_frequency_max_percent":float(np.max(positive_values)) if len(positive_values) else None,
            "populations_positive_at_least_2_years":int((g.positive_years>=2).sum()),
            "populations_positive_at_least_3_years":int((g.positive_years>=3).sum()),
        }

    traj=pd.DataFrame(trajectory_rows)
    traj.to_csv(out/"population_trajectories.csv",index=False)

    pal_pos=(traj.loc[(traj.phenotype=="PAL") & (traj.positive_years>0)]
               .sort_values(["positive_years","max_frequency"],ascending=[False,False]))
    top_pal=[]
    for _,r in pal_pos.iterrows():
        pop=r.population
        q=long[(long.population==pop)&(long.phenotype=="PAL")].sort_values("year")
        top_pal.append({
            "population":pop,
            "observed_years":[int(x) for x in q.year],
            "frequencies_percent":[float(x) for x in q.frequency_percent],
            "positive_years":int((q.frequency_percent>0).sum()),
            "range_percent":[float(q.frequency_percent.min()),float(q.frequency_percent.max())],
        })

    result={
        "schema":"fcp_silene_decoupling_natural_persistence_v1",
        "status":"complete",
        "source":"Supplementary Table S2, Del Valle et al. 2019, DOI 10.1186/s12870-019-2082-6",
        "surveyed_populations":int(len(pops)),
        "phenotypes":phenotype_results,
        "persistent_PAL_populations":top_pal,
        "descriptive_contrast":{
            "PAL_positive_frequency_median_percent":phenotype_results["PAL"]["positive_frequency_median_percent"],
            "WAL_positive_frequency_median_percent":phenotype_results["WAL"]["positive_frequency_median_percent"],
            "PAL_over_WAL_positive_median_ratio":float(
                phenotype_results["PAL"]["positive_frequency_median_percent"]/
                phenotype_results["WAL"]["positive_frequency_median_percent"]
            ),
            "PAL_positive_frequency_min_percent":phenotype_results["PAL"]["positive_frequency_min_percent"],
            "WAL_positive_frequency_max_percent":phenotype_results["WAL"]["positive_frequency_max_percent"],
            "all_positive_PAL_frequencies_exceed_max_positive_WAL":bool(
                phenotype_results["PAL"]["positive_frequency_min_percent"] >
                phenotype_results["WAL"]["positive_frequency_max_percent"]
            ),
        },
        "interpretation":[
            "PAL is geographically restricted but reaches moderate frequency and is repeatedly observed in the same populations.",
            "WAL occurs in more populations but remains below 1% in every positive population-year in Table S2.",
            "Within this species, petal-restricted anthocyanin loss is therefore associated with much greater local frequency/persistence than whole-plant anthocyanin loss."
        ],
        "hard_nonclaims":[
            "post-publication reproduction, not independent confirmation",
            "does not prove tissue restriction itself causally causes higher PAL frequency",
            "natural frequency can also reflect demography, mutation rate, selection and spatial history"
        ]
    }
    (out/"result.json").write_text(json.dumps(result,indent=2)+"\n")
    print(json.dumps(result,indent=2))

if __name__=="__main__":
    main()
