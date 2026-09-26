#!/usr/bin/env python3
from __future__ import annotations
import argparse,json,re
from pathlib import Path
import numpy as np
import pandas as pd
from docx import Document

RANGE_RX=re.compile(r"^\s*([0-9.]+)\s*[-–]\s*([0-9.]+)")
LT_RX=re.compile(r"^\s*<\s*([0-9.]+)")
NUM_RX=re.compile(r"^\s*([0-9.]+)\s*(?:\(|$)")

def family_from_species_cell(s):
    m=re.search(r"\(([^()]*)\)\s*$",s)
    return m.group(1).strip() if m else ""

def species_name(s):
    return re.sub(r"\s*\([^()]*\)\s*$","",s).strip()

def parse_freq(s,cls):
    raw=str(s).strip().replace("“","").replace("”","")
    out={"raw_frequency":raw,"lower_bound_percent":np.nan,"upper_bound_percent":np.nan,
         "qualitative":"","censoring":"unparsed","parenthetical_mean_percent":np.nan}
    if cls=="WAL" and "greenhouse" in raw.lower():
        out["qualitative"]="greenhouse"
        out["censoring"]="excluded_greenhouse"
        m=re.search(r"([0-9.]+)",raw)
        if m:
            out["lower_bound_percent"]=float(m.group(1))
            out["upper_bound_percent"]=float(m.group(1))
        return out
    m=RANGE_RX.search(raw)
    if m:
        out["lower_bound_percent"]=float(m.group(1)); out["upper_bound_percent"]=float(m.group(2))
        out["censoring"]="range"
        mm=re.search(r"mean\s*([0-9.]+)",raw,re.I)
        if mm: out["parenthetical_mean_percent"]=float(mm.group(1))
        return out
    m=LT_RX.search(raw)
    if m:
        out["lower_bound_percent"]=0.0; out["upper_bound_percent"]=float(m.group(1))
        out["censoring"]="left_censored_upper_bound"
        return out
    m=NUM_RX.search(raw)
    if m:
        v=float(m.group(1)); out["lower_bound_percent"]=v; out["upper_bound_percent"]=v
        out["censoring"]="exact"
        return out
    low=raw.lower()
    if "extremely rare" in low:
        out["qualitative"]="extremely rare"; out["censoring"]="qualitative"
    elif "rare" in low:
        out["qualitative"]="rare"; out["censoring"]="qualitative"
    return out

def main():
    p=argparse.ArgumentParser()
    p.add_argument("--docx",required=True); p.add_argument("--outdir",required=True)
    a=p.parse_args()
    out=Path(a.outdir); out.mkdir(parents=True,exist_ok=True)
    doc=Document(a.docx)
    if len(doc.tables)!=1: raise SystemExit("expected one table")
    t=doc.tables[0]
    rows=[]
    cls=None
    for row in t.rows[3:]:
        c=[x.text.strip() for x in row.cells]
        if not c or not c[0]: continue
        if c[0].startswith("Species with PAL"):
            cls="PAL"; continue
        if c[0].startswith("Species with WAL"):
            cls="WAL"; continue
        if c[0].startswith("a, all polymorphic"):
            break
        if cls not in {"PAL","WAL"}: continue
        rec={"phenotype_class":cls,"species_source":c[0],"species":species_name(c[0]),
             "family_source":family_from_species_cell(c[0]),"reference":c[4],
             "nonanthocyanin_flavonoids_flower":c[2],
             "nonanthocyanin_flavonoids_vegetative":c[3]}
        rec.update(parse_freq(c[1],cls))
        rows.append(rec)
    d=pd.DataFrame(rows)
    d.to_csv(out/"source_frequency_registry.csv",index=False)

    pal=d[d.phenotype_class.eq("PAL")].copy()
    wal=d[d.phenotype_class.eq("WAL")].copy()
    wal_num=wal[wal.censoring.isin(["exact","left_censored_upper_bound"])].copy()
    wal_qual=wal[wal.censoring.eq("qualitative")].copy()
    wal_nat=wal[~wal.censoring.eq("excluded_greenhouse")].copy()

    max_wal=float(wal_num.upper_bound_percent.max())
    lower_pal=pal.lower_bound_percent.dropna().to_numpy(float)
    upper_pal=pal.upper_bound_percent.dropna().to_numpy(float)
    upper_wal=wal_num.upper_bound_percent.dropna().to_numpy(float)

    result={
      "schema":"fcp_crossspecies_pal_wal_frequency_v1",
      "status":"complete",
      "source":"Supplementary Table S1, Del Valle et al. 2019, DOI 10.1186/s12870-019-2082-6",
      "PAL":{
        "instances":int(len(pal)),
        "source_taxonomic_groups":int(pal.family_source.nunique()),
        "lower_bound_median_percent":float(np.median(lower_pal)),
        "upper_bound_median_percent":float(np.median(upper_pal)),
        "minimum_upper_endpoint_percent":float(np.min(upper_pal)),
        "maximum_upper_endpoint_percent":float(np.max(upper_pal)),
        "instances_lower_bound_above_max_quantified_WAL":int((lower_pal>max_wal).sum())
      },
      "WAL":{
        "instances_total":int(len(wal)),
        "instances_natural_non_greenhouse":int(len(wal_nat)),
        "source_taxonomic_groups":int(wal.family_source.nunique()),
        "numeric_natural_instances":int(len(wal_num)),
        "numeric_upper_bound_median_percent":float(np.median(upper_wal)),
        "numeric_upper_bound_max_percent":max_wal,
        "qualitative_rare_instances":int((wal_qual.qualitative=="rare").sum()),
        "qualitative_extremely_rare_instances":int((wal_qual.qualitative=="extremely rare").sum()),
        "excluded_greenhouse_instances":int((wal.censoring=="excluded_greenhouse").sum())
      },
      "descriptive_contrast":{
        "PAL_median_lower_bound_over_WAL_median_upper_bound":float(np.median(lower_pal)/np.median(upper_wal)),
        "PAL_minimum_upper_endpoint_over_WAL_maximum_upper_bound":float(np.min(upper_pal)/max_wal),
        "all_PAL_upper_endpoints_above_max_quantified_WAL":bool(np.all(upper_pal>max_wal))
      },
      "interpretation":[
        "PAL systems commonly reach multi-percent to high-frequency polymorphism.",
        "Quantified natural WAL systems are bounded at or below low single-percent frequency; additional WAL rows are described qualitatively as rare or extremely rare.",
        "The literature synthesis is consistent with higher persistence/frequency when anthocyanin loss is restricted to petals rather than extending through the whole plant."
      ],
      "hard_nonclaims":[
        "post-publication literature synthesis, not an unbiased species sample",
        "does not prove tissue restriction causally causes higher frequency",
        "source family labels are preserved as printed and not taxonomically corrected",
        "qualitative WAL frequencies are not numerically imputed"
      ]
    }
    (out/"result.json").write_text(json.dumps(result,indent=2)+"\n")
    print(json.dumps(result,indent=2))

if __name__=="__main__": main()
