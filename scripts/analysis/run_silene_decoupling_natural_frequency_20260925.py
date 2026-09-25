#!/usr/bin/env python3
from __future__ import annotations
import argparse, json, re
from pathlib import Path
import numpy as np
import pandas as pd
from docx import Document

YEARS=[2010,2013,2014,2016,2018]

def parse_num(x):
    s=str(x).strip()
    if s in {"","-","–","—"}:
        return np.nan
    try:
        return float(s)
    except Exception:
        return np.nan

def main():
    p=argparse.ArgumentParser()
    p.add_argument("--docx",required=True)
    p.add_argument("--outdir",required=True)
    a=p.parse_args()
    out=Path(a.outdir); out.mkdir(parents=True,exist_ok=True)

    doc=Document(a.docx)
    if len(doc.tables)!=1:
        raise SystemExit(f"expected exactly one table, found {len(doc.tables)}")
    t=doc.tables[0]
    if len(t.columns)!=14 or len(t.rows)<24:
        raise SystemExit(f"unexpected table shape rows={len(t.rows)} cols={len(t.columns)}")

    rows=[[c.text.replace("\n"," | ").strip() for c in r.cells] for r in t.rows]
    h1=rows[1]; h2=rows[2]
    expected_years=[str(y) for y in YEARS for _ in (0,1)]
    if h1[0]!="Locality" or h1[1]!="Latitude" or h1[2]!="Longitude":
        raise SystemExit(f"unexpected primary header {h1[:4]}")
    if h1[4:14]!=expected_years:
        raise SystemExit(f"unexpected year header {h1[4:14]}")
    if h2[4:14]!=["PAL","WAL"]*5:
        raise SystemExit(f"unexpected PAL/WAL header {h2[4:14]}")

    recs=[]; populations=[]
    for source_row,cells in enumerate(rows[3:],start=4):
        locality=cells[0].strip()
        if not locality or "represents the absence of information" in locality:
            continue
        populations.append({
            "locality":locality,
            "latitude":cells[1],
            "longitude":cells[2],
            "estimated_population_size":cells[3],
            "source_excel_row":source_row,
        })
        for k,year in enumerate(YEARS):
            pal=parse_num(cells[4+2*k])
            wal=parse_num(cells[5+2*k])
            recs.append({
                "locality":locality,
                "year":year,
                "PAL_percent":pal,
                "WAL_percent":wal,
                "PAL_observed":bool(np.isfinite(pal)),
                "WAL_observed":bool(np.isfinite(wal)),
                "source_excel_row":source_row,
            })

    pop=pd.DataFrame(populations).drop_duplicates("locality")
    long=pd.DataFrame(recs).sort_values(["locality","year"])
    pop.to_csv(out/"population_metadata.csv",index=False)
    long.to_csv(out/"population_year_frequencies.csv",index=False)

    phenotype={}
    traj=[]
    for short,col in [("PAL","PAL_percent"),("WAL","WAL_percent")]:
        q=long.loc[long[col].notna()].copy()
        pos=q.loc[q[col]>0].copy()
        by=pos.groupby("locality").agg(
            positive_years=("year","nunique"),
            min_positive_percent=(col,"min"),
            median_positive_percent=(col,"median"),
            max_positive_percent=(col,"max"),
            first_positive_year=("year","min"),
            last_positive_year=("year","max"),
        ).reset_index()
        by["phenotype"]=short
        traj.append(by)
        phenotype[short]={
            "population_years_observed":int(len(q)),
            "population_years_positive":int(len(pos)),
            "populations_ever_positive":int(pos.locality.nunique()),
            "positive_frequency_min_percent":float(pos[col].min()) if len(pos) else None,
            "positive_frequency_median_percent":float(pos[col].median()) if len(pos) else None,
            "positive_frequency_max_percent":float(pos[col].max()) if len(pos) else None,
            "populations_positive_ge_2_years":int((by.positive_years>=2).sum()) if len(by) else 0,
            "populations_positive_ge_3_years":int((by.positive_years>=3).sum()) if len(by) else 0,
        }

    trajectories=pd.concat(traj,ignore_index=True)
    trajectories.to_csv(out/"positive_frequency_trajectories.csv",index=False)

    paltraj=trajectories.loc[trajectories.phenotype.eq("PAL")].sort_values(
        ["positive_years","max_positive_percent"],ascending=[False,False]
    )
    longest_pal=[]
    for r in paltraj.head(2).itertuples(index=False):
        vals=long.loc[(long.locality.eq(r.locality)) & long.PAL_percent.notna(),
                      ["year","PAL_percent"]]
        longest_pal.append({
            "locality":r.locality,
            "positive_years":int(r.positive_years),
            "observed_positive_range_percent":[float(r.min_positive_percent),float(r.max_positive_percent)],
            "trajectory":[{"year":int(x.year),"PAL_percent":float(x.PAL_percent)}
                          for x in vals.itertuples(index=False)]
        })

    reproduction=bool(
        phenotype["PAL"]["populations_ever_positive"]==2 and
        phenotype["PAL"]["positive_frequency_min_percent"]>=8 and
        phenotype["PAL"]["positive_frequency_max_percent"]<=21 and
        phenotype["PAL"]["populations_positive_ge_3_years"]==2 and
        phenotype["WAL"]["populations_ever_positive"]==9 and
        phenotype["WAL"]["positive_frequency_max_percent"]<1
    )

    result={
        "schema":"fcp_silene_decoupling_natural_frequency_v1",
        "status":"complete",
        "role":"post_publication_public_table_reproduction",
        "source":"Del Valle et al. 2019 BMC Plant Biology 19:496; Supplementary Table S2",
        "doi":"10.1186/s12870-019-2082-6",
        "surveyed_populations":int(pop.locality.nunique()),
        "survey_years":YEARS,
        "phenotypes":phenotype,
        "longest_repeated_PAL_series":longest_pal,
        "published_pattern_reproduced":reproduction,
        "interpretation":[
            "PAL is geographically restricted to two populations but persists at moderate frequency across repeated survey years.",
            "WAL occurs in more populations but only at very low frequency, with all positive observations below 1%.",
            "The pattern concerns maintenance/frequency rather than mutational occurrence: whole-plant loss can appear repeatedly yet remain rare."
        ],
        "hard_nonclaims":[
            "post-publication reproduction, not independent confirmation",
            "does not prove tissue restriction is the sole cause of PAL persistence",
            "rounded percentages are not converted into inferred individual counts",
            "does not alter the frozen New Phytologist manuscript"
        ]
    }
    (out/"result.json").write_text(json.dumps(result,indent=2,ensure_ascii=False)+"\n")
    print(json.dumps(result,indent=2,ensure_ascii=False))

if __name__=="__main__":
    main()
