#!/usr/bin/env python3
"""Build a review-ready New Phytologist DOCX from the frozen Markdown manuscript.

This is a packaging transform only. It does not change scientific claims.
Administrative placeholders are deliberately preserved until authors provide
final metadata.
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


FIGURE_FILES = {
    1: "polymorphism_figure1_measurement_frame.png",
    2: "polymorphism_figure2_h1_reproducibility.png",
    3: "polymorphism_figure3_h2_target_localization.png",
    4: "polymorphism_figure4_prospective_h2.png",
    5: "polymorphism_figure5_explanatory_boundaries.png",
}


def _set_run_font(run, name: str = "Times New Roman", size: float | None = 12) -> None:
    run.font.name = name
    if size is not None:
        run.font.size = Pt(size)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), name)


def _configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)

    for style_name, size in [
        ("Title", 16),
        ("Heading 1", 14),
        ("Heading 2", 13),
        ("Heading 3", 12),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.paragraph_format.line_spacing = 1.5
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(4)

    for style_name in ("List Bullet", "List Number"):
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(12)
        style.paragraph_format.line_spacing = 1.5


def _set_section_layout(section) -> None:
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.15)
    section.right_margin = Inches(1.0)

    sect_pr = section._sectPr
    for node in list(sect_pr):
        if node.tag == qn("w:lnNumType"):
            sect_pr.remove(node)
    line_numbers = OxmlElement("w:lnNumType")
    line_numbers.set(qn("w:countBy"), "1")
    line_numbers.set(qn("w:distance"), "360")
    line_numbers.set(qn("w:restart"), "continuous")
    sect_pr.append(line_numbers)

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.text = ""
    run = p.add_run("Page ")
    _set_run_font(run, size=10)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def _add_inline_markdown(paragraph, text: str, *, base_size: float = 12) -> None:
    """Add a small safe subset of Markdown inline formatting."""
    token_re = re.compile(r"(\*\*.+?\*\*|(?<!\*)\*[^*]+?\*(?!\*)|\`[^\`]+?\`)")
    pos = 0
    for match in token_re.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos : match.start()])
            _set_run_font(run, size=base_size)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
            _set_run_font(run, size=base_size)
        elif token.startswith("*"):
            run = paragraph.add_run(token[1:-1])
            run.italic = True
            _set_run_font(run, size=base_size)
        else:
            run = paragraph.add_run(token[1:-1])
            _set_run_font(run, name="Courier New", size=base_size - 1)
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        _set_run_font(run, size=base_size)


def _is_table_separator(line: str) -> bool:
    cells = [c.strip() for c in line.strip().strip("|").split("|")]
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", c or "") for c in cells)


def _parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    i = start
    while i < len(lines) and lines[i].lstrip().startswith("|"):
        if not _is_table_separator(lines[i]):
            rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
        i += 1
    return rows, i


def _add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = "Table Grid"
    table.autofit = True
    for r_idx, row in enumerate(rows):
        for c_idx in range(ncols):
            cell = table.cell(r_idx, c_idx)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(0)
            txt = row[c_idx] if c_idx < len(row) else ""
            _add_inline_markdown(p, txt, base_size=9.5)
            if r_idx == 0:
                for run in p.runs:
                    run.bold = True
    doc.add_paragraph("")


def _add_figure(doc: Document, figure_number: int, figures_dir: Path) -> None:
    path = figures_dir / FIGURE_FILES[figure_number]
    if not path.exists():
        raise FileNotFoundError(path)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_together = True
    run = p.add_run()
    run.add_picture(str(path), width=Inches(6.15))


def _add_formula_paragraph(doc: Document, formula: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_together = True
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(formula)
    _set_run_font(run, name="Cambria Math", size=11)


def _normalize_formula(text: str) -> str:
    return (
        text.replace("operatorname{normalize}", "normalize")
        .replace("mathrm{", "")
        .replace("}", "")
        .replace("sum_k", "Σ_k")
    )


def _add_body_paragraph(doc: Document, text: str, style: str | None = None) -> None:
    p = doc.add_paragraph(style=style)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.widow_control = True
    if style == "List Bullet":
        p.paragraph_format.space_after = Pt(2)
    _add_inline_markdown(p, text)


def _add_title_page_note(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    run = p.add_run("Review-ready submission formatting; unresolved administrative placeholders are intentionally retained.")
    run.italic = True
    _set_run_font(run, size=9)


def build_submission_docx(
    source_path: Path,
    figures_dir: Path,
    output_path: Path,
) -> Path:
    source_path = Path(source_path)
    figures_dir = Path(figures_dir)
    output_path = Path(output_path)

    if not source_path.exists():
        raise FileNotFoundError(source_path)
    for filename in FIGURE_FILES.values():
        if not (figures_dir / filename).exists():
            raise FileNotFoundError(figures_dir / filename)

    lines = source_path.read_text(encoding="utf-8").splitlines()
    doc = Document()
    _configure_styles(doc)
    for section in doc.sections:
        _set_section_layout(section)

    in_figure_legends = False
    added_title_note = False
    i = 0
    while i < len(lines):
        line = lines[i].strip()

        if not line or line == "---":
            i += 1
            continue

        if line.startswith("|"):
            rows, i = _parse_table(lines, i)
            _add_table(doc, rows)
            continue

        if line.startswith("# "):
            p = doc.add_paragraph(style="Title")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _add_inline_markdown(p, line[2:].strip(), base_size=16)
            if not added_title_note:
                _add_title_page_note(doc)
                added_title_note = True
            i += 1
            continue

        if line.startswith("## "):
            heading = line[3:].strip()
            doc.add_paragraph(heading, style="Heading 1")
            if heading == "Figure legends":
                in_figure_legends = True
            elif heading == "Supporting Information":
                in_figure_legends = False
            i += 1
            continue

        if line.startswith("### "):
            doc.add_paragraph(line[4:].strip(), style="Heading 2")
            i += 1
            continue

        if line.startswith("#### "):
            doc.add_paragraph(line[5:].strip(), style="Heading 3")
            i += 1
            continue

        if line == "[" and i + 2 < len(lines):
            formula = lines[i + 1].strip()
            if lines[i + 2].strip() == "]":
                _add_formula_paragraph(doc, _normalize_formula(formula))
                i += 3
                continue

        if line.startswith("- "):
            _add_body_paragraph(doc, line[2:].strip(), style="List Bullet")
            i += 1
            continue

        if re.match(r"^\d+\.\s+", line):
            _add_body_paragraph(doc, re.sub(r"^\d+\.\s+", "", line), style="List Number")
            i += 1
            continue

        figure_match = re.match(r"^\*\*Figure\s+(\d+)\.", line)
        if in_figure_legends and figure_match:
            figure_no = int(figure_match.group(1))
            if figure_no > 1:
                doc.add_page_break()
            p = doc.add_paragraph()
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.line_spacing = 1.5
            _add_inline_markdown(p, line)
            if figure_no in FIGURE_FILES:
                _add_figure(doc, figure_no, figures_dir)
            i += 1
            continue

        paragraph_lines = [line]
        j = i + 1
        while j < len(lines):
            nxt = lines[j].strip()
            if (
                not nxt
                or nxt == "---"
                or nxt.startswith("#")
                or nxt.startswith("- ")
                or nxt.startswith("|")
                or nxt == "["
                or re.match(r"^\d+\.\s+", nxt)
                or (in_figure_legends and re.match(r"^\*\*Figure\s+\d+\.", nxt))
            ):
                break
            paragraph_lines.append(nxt)
            j += 1
        _add_body_paragraph(doc, " ".join(paragraph_lines))
        i = j

    for section in doc.sections:
        _set_section_layout(section)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    return output_path


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--figures", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    out = build_submission_docx(args.source, args.figures, args.output)
    print(out)


if __name__ == "__main__":
    main()
